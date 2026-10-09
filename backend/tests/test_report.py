"""Tests of the report: grading, the AI answer check, the self-check, and a whole run with
the database steps replaced by fakes. No database or AI is used.

Run from backend/:  python -m pytest tests -q
"""
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

from app.report import jobs, plain, settings, synth, testmode, validate


def evidence(results1=("Correct",), results2=("Must fix",)):
    """A small evidence in the shape collect.build makes."""
    def p1(n, r):
        return {"part": 1, "id": f"1|1|dbo.t{n}", "source_server": "ATNM", "source_db": "AppDb", "target_server": "RDS",
                "target_db": "AppDb_Copy", "table": f"dbo.t{n}", "mapping": f"m{n}", "in_source": True,
                "in_target": True, "rows_source": 10, "rows_target": 10 if r == "Correct" else 9,
                "rows_diff": 0 if r == "Correct" else -1, "columns_source": 3, "columns_target": 3,
                "columns_missing": 0, "columns_extra": 0, "columns_changed": 0, "columns_renamed": 0,
                "keys_rules": "all kept", "values": "identical", "values_detail": "", "checked_at": "2026-10-03T10:00:00+00:00",
                "live": False, "result": r, "reason": "why"}

    def p2(n, r):
        return {"part": 2, "id": f"2|m{n}", "mapping": f"m{n}", "type": "one_to_one", "source_db": "AppDb_Copy",
                "target_db": "NewDb", "source": f"dbo.t{n}", "target": f"new.t{n}", "rule": "target rows = source rows",
                "rows_expected": 10, "rows_actual": 10, "rows_diff": 0, "row_count": "Correct", "columns": "Correct",
                "values": r, "key_links": "Correct", "renames": "0 confirmed", "target_health": "Correct",
                "checked_at": "2026-10-03T10:00:00+00:00", "result": r, "reason": "why"}

    items1 = [p1(n, r) for n, r in enumerate(results1)]
    items2 = [p2(n, r) for n, r in enumerate(results2)]
    issues = []
    for it in items1 + items2:
        if it["result"] in (plain.MUST_FIX, plain.DECIDE):
            issues.append({"id": f"{it['part']}-{len(issues) + 1:03d}", "part": it["part"], "item": it["id"],
                           "source_db": it["source_db"], "target_db": it["target_db"],
                           "table": it.get("table") or it.get("source"), "target_table": it.get("target") or it.get("table"),
                           "mapping": it["mapping"], "column": "", "check": "Values", "result": it["result"],
                           "finding": "3 records differ.", "meaning": plain.meaning("Values"),
                           "action": plain.action("Values"), "checked_at": it["checked_at"]})

    def tally(items):
        out = {r: 0 for r in plain.RESULTS}
        for x in items:
            out[x["result"]] += 1
        return out

    return {
        "run": {"id": "T1", "started_at": "2026-10-03T09:00:00+00:00", "finished_at": "2026-10-03T10:30:00+00:00",
                "mode": "test"},
        "overall": plain.overall(x["result"] for x in items1 + items2),
        "totals": {"part1": tally(items1), "part2": tally(items2),
                   "issues": {r: sum(1 for i in issues if i["result"] == r) for r in plain.RESULTS}},
        "part1": {"items": items1, "columns": [], "constraints": [], "empty": [], "renames": [], "row_diffs": [],
                  "identity": [], "databases": [{"pair": "1", "source_server": "ATNM", "source_db": "AppDb",
                                                 "target_server": "RDS", "target_db": "AppDb_Copy",
                                                 "required": len(items1)}], "errors": []},
        "part2": {"items": items2, "columns": [], "renames": [], "values": [], "links": [], "health": [], "identity": [],
                  "databases": [{"side": "T", "name": "NewDb", "role": "target"}]},
        "issues": issues, "not_checked": [], "coverage": {"databases": []},
    }


class Grading(unittest.TestCase):
    def test_overall(self):
        self.assertEqual(plain.overall([plain.CORRECT, plain.MUST_FIX, plain.NOT_CHECKED]), "NOT READY")
        self.assertEqual(plain.overall([plain.CORRECT, plain.NOT_CHECKED, plain.DECIDE]), "INCOMPLETE")
        self.assertEqual(plain.overall([plain.CORRECT, plain.DECIDE, plain.BY_DESIGN]), "READY")

    def test_worst(self):
        self.assertEqual(plain.worst(plain.CORRECT, plain.DECIDE, plain.MUST_FIX), plain.MUST_FIX)
        self.assertEqual(plain.worst(plain.CORRECT, plain.NOT_CHECKED, plain.DECIDE), plain.NOT_CHECKED)
        self.assertEqual(plain.worst(plain.BY_DESIGN, plain.BY_DESIGN), plain.BY_DESIGN)

    def test_every_kind_is_explained(self):
        for kind in plain.KINDS:
            self.assertTrue(plain.meaning(kind) and plain.action(kind))


class DataOnly(unittest.TestCase):
    """The report grades the data: table rules and column settings other than the type are left out."""

    def test_rules_and_settings_are_not_graded(self):
        from app.ATNM import structure
        from app.report import collect
        soft_col = {"name": "Qty", "notes": ["may now be empty"], "severity": "review"}
        soft_text = structure._names([f"{soft_col['name']}: {', '.join(soft_col['notes'])}"], 3)
        t = {"columns": {"changed": [soft_col], "renamed": []}, "reasons": [
            {"severity": "review", "text": "8 keys or rules not kept in RDS: 7 default values, 1 check rule"},
            {"severity": "problem", "text": "Primary key (UId) not in RDS"},
            {"severity": "problem", "text": "Identity counter of Id in RDS is behind (next value 5, 9 already used)"},
            {"severity": "review", "text": soft_text},
            {"severity": "problem", "text": "RDS has 3 fewer rows (123 in ATNM, 120 in RDS)"},
            {"severity": "problem", "text": "1 column type changed: Qty (type int → bigint)"},
        ]}
        kept = [r["text"] for r in collect._data_reasons(t, structure)]
        self.assertEqual(kept, ["RDS has 3 fewer rows (123 in ATNM, 120 in RDS)",
                                "1 column type changed: Qty (type int → bigint)"])

    def test_a_type_change_is_kept_even_among_settings(self):
        from app.ATNM import structure
        from app.report import collect
        col = {"name": "Qty", "notes": ["type int → bigint", "may now be empty"], "severity": "review"}
        text = structure._names([f"{col['name']}: {', '.join(col['notes'])}"], 3)
        t = {"columns": {"changed": [col], "renamed": []}, "reasons": [{"severity": "review", "text": text}]}
        self.assertEqual(len(collect._data_reasons(t, structure)), 1)


class AiAnswerCheck(unittest.TestCase):
    FACTS = "part1: 65 required tables: 41 Correct\n- part1 | Must fix | Row count | AppMasterDB.dbo.task_daily | 320 fewer"

    def answer(self, text):
        return {"executive_summary": text, "part1_summary": "", "part2_summary": "", "top_risks": [], "next_steps": []}

    def test_numbers_and_names_from_the_facts_pass(self):
        self.assertEqual(synth._bad(self.answer("41 of 65 tables are correct; dbo.task_daily has 320 fewer records."),
                                    self.FACTS), [])

    def test_invented_number_is_caught(self):
        self.assertIn("999", synth._bad(self.answer("999 records are missing."), self.FACTS))

    def test_invented_table_is_caught(self):
        self.assertIn("dbo.payroll", synth._bad(self.answer("dbo.payroll is broken."), self.FACTS))

    def test_template_needs_no_ai(self):
        t = synth.template(evidence())
        self.assertTrue(t["executive_summary"] and t["part1_summary"] and t["part2_summary"])


class SelfCheck(unittest.TestCase):
    def test_consistent_evidence(self):
        self.assertEqual(validate.evidence(evidence(), 1, 1), [])

    def test_missing_table_is_reported(self):
        self.assertTrue(any("requires 2" in p for p in validate.evidence(evidence(), 2, 1)))

    def test_wrong_overall_is_reported(self):
        ev = evidence()
        ev["overall"] = "READY"
        self.assertTrue(any("overall" in p for p in validate.evidence(ev, 1, 1)))

    def test_files_agree_with_evidence(self):
        ev = evidence(("Correct", "Must fix"), ("Needs a decision", "Correct", "By design"))
        ev["summary"] = {"source": "template", **synth.template(ev)}
        with tempfile.TemporaryDirectory() as d:
            x, w = Path(d) / "r.xlsx", Path(d) / "r.docx"
            from app.report import docx_writer, xlsx_writer
            xlsx_writer.build(ev, x)
            docx_writer.build(ev, w)
            self.assertEqual(validate.files(ev, x, w), [])

    def test_two_sections_in_both_files(self):
        ev = evidence(("Correct", "Must fix"), ("Needs a decision", "Correct"))
        ev["summary"] = {"source": "template", **synth.template(ev)}
        with tempfile.TemporaryDirectory() as d:
            x, w = Path(d) / "r.xlsx", Path(d) / "r.docx"
            from docx import Document
            from openpyxl import load_workbook

            from app.report import docx_writer, xlsx_writer
            xlsx_writer.build(ev, x)
            docx_writer.build(ev, w)
            wb = load_workbook(x)
            # Summary, then the ATNM sheet, then the RDS sheet; the details after them
            self.assertEqual(wb.sheetnames[:3], ["Summary", xlsx_writer.P1_TABLES, xlsx_writer.P2_MAPPINGS])
            ws = wb[xlsx_writer.P1_TABLES]
            head = [c.value for c in ws[xlsx_writer.FIRST - 1]]
            self.assertEqual(head[:5], ["Source table (ATNM)", "Row count", "Target table (RDS)", "Row count", "Difference"])
            self.assertEqual(head[5:8], ["Columns", "Renamed columns", "Null values"])
            self.assertEqual(head[8:10], ["Status", "Issues found"])
            self.assertNotIn("Used by mapping", head)      # the copy is table for table
            rows = {r[0]: r for r in ws.iter_rows(min_row=xlsx_writer.FIRST, values_only=True)}
            # the problem row lists its issue with the same number as the Word report; the correct one says so
            bad = rows["AppDb · dbo.t1"]
            self.assertEqual(bad[8], "Need to fix")      # the workbook's word for Must fix
            self.assertIn(next(i["id"] for i in ev["issues"] if i["part"] == 1), bad[9])
            self.assertTrue(rows["AppDb · dbo.t0"][9].startswith("No issue"))
            # the two main sheets, the two kept ones and Full Details besides the Summary
            self.assertEqual(wb.sheetnames[3:], [xlsx_writer.NOT_IN_PLAN, xlsx_writer.WORDS, xlsx_writer.FULL])
            # the Summary is one paragraph; Full Details has no per-table list
            self.assertTrue(wb["Summary"]["A5"].value.startswith(plain.SUMMARY_INTRO))
            self.assertIn("1. About this report", [c.value for c in wb[xlsx_writer.FULL]["A"]])
            # every cell of a list has a border
            self.assertTrue(all(c.border.left.style and c.border.top.style for c in ws[xlsx_writer.FIRST]))
            ws2 = wb[xlsx_writer.P2_MAPPINGS]
            self.assertEqual([c.value for c in ws2[xlsx_writer.FIRST - 1]][:2], ["Mapping", "Source table(s) (RDS)"])
            heads = [p.text for p in Document(w).paragraphs if p.style.name.startswith("Heading 1")]
            s1 = next(i for i, h in enumerate(heads) if h.startswith("Section 1"))
            s2 = next(i for i, h in enumerate(heads) if h.startswith("Section 2"))
            self.assertLess(s1, s2)
            self.assertIn("ATNM → RDS", heads[s1])


class WholeRun(unittest.TestCase):
    """jobs.start with Part 1 and Part 2 faked: stages, progress, files and self-check."""

    def setUp(self):
        # The normal report, whatever REPORT_TEST_TABLES the .env sets.
        p = mock.patch.object(settings, "TEST_TABLES", 0)
        p.start()
        self.addCleanup(p.stop)

    def test_run_writes_both_files(self):
        with tempfile.TemporaryDirectory() as d:
            ev = evidence(("Correct", "Correct"), ("Correct", "Must fix"))
            atnm_state = {"running": False, "done": 2, "total": 2, "results": {"identical": 2}}
            fake_main = mock.MagicMock()
            fake_main._checks.run_all.return_value = True
            fake_main._checks.job = {"running": False, "done": 2, "total": 2, "current": None}
            fake_main.get_plan.return_value.mappings = [mock.Mock(id="m0"), mock.Mock(id="m1")]
            pf = {"ok": True, "atnm": {"ok": True}, "rds": {"ok": True}, "busy": [],
                  "plan": {"ok": True, "tables": 2, "mappings": 2}, "ai": {"enabled": False}}
            with mock.patch.object(settings, "REPORT_DIR", Path(d)), \
                    mock.patch.object(settings, "RUN_FILE", Path(d) / "run.json"), \
                    mock.patch.object(jobs, "preflight", return_value=pf), \
                    mock.patch.object(jobs, "_part1", side_effect=lambda r, s: jobs._stage("part1", "done", done=2, total=2)), \
                    mock.patch.object(jobs, "_part2", side_effect=lambda r, s: jobs._stage("part2", "done", done=2, total=2)), \
                    mock.patch("app.report.collect.build", return_value=ev), \
                    mock.patch.object(synth, "write", return_value={"source": "template", **synth.template(ev)}):
                ok, msg = jobs.start()
                self.assertTrue(ok, msg)
                for _ in range(100):
                    if not jobs.job["running"]:
                        break
                    time.sleep(0.1)
                st = jobs.status(0)
                self.assertFalse(st["running"])
                self.assertIsNone(st["error"], st["error"])
                self.assertEqual([s["status"] for s in st["stages"]], ["done"] * 6)
                self.assertEqual(st["fraction"], 1.0)
                meta = st["result"]
                self.assertEqual(meta["overall"], "NOT READY")
                self.assertEqual(meta["self_check"], [])
                self.assertTrue(jobs.file_of(meta["run_id"], "docx").exists())
                self.assertTrue(jobs.file_of(meta["run_id"], "xlsx").exists())
                self.assertIsNone(jobs.file_of(meta["run_id"], "../../etc"))
                self.assertEqual(len(jobs.reports()), 1)

    def test_refused_when_preflight_fails(self):
        pf = {"ok": False, "atnm": {"ok": False, "label": "ATNM", "error": "Connect the VPN."}, "rds": {"ok": True},
              "busy": [], "plan": {"ok": True, "tables": 2, "mappings": 2}, "ai": {"enabled": False}}
        with mock.patch.object(jobs, "preflight", return_value=pf):
            ok, msg = jobs.start()
        self.assertFalse(ok)
        self.assertIn("VPN", msg)


class WhoMayStart(unittest.TestCase):
    def request(self, host, forwarded_host=None, client="127.0.0.1"):
        headers = {"host": host}
        if forwarded_host:
            headers["x-forwarded-host"] = forwarded_host
        return mock.Mock(headers=headers, client=mock.Mock(host=client))

    def test_this_computer_via_localhost(self):
        from app.report import api
        with mock.patch.object(settings, "ALLOW_REMOTE", False):
            self.assertTrue(api._allowed(self.request("127.0.0.1:8000", "localhost:3000")))
            self.assertTrue(api._allowed(self.request("127.0.0.1:8000")))           # straight to the backend

    def test_colleague_via_wifi_address(self):
        from app.report import api
        with mock.patch.object(settings, "ALLOW_REMOTE", False):
            self.assertFalse(api._allowed(self.request("127.0.0.1:8000", "192.168.1.10:3000")))

    def test_allow_remote_setting(self):
        from app.report import api
        with mock.patch.object(settings, "ALLOW_REMOTE", True):
            self.assertTrue(api._allowed(self.request("127.0.0.1:8000", "192.168.1.10:3000")))


class TestMode(unittest.TestCase):
    """REPORT_TEST_TABLES: the smallest tables with records, both parts, files marked TEST."""

    def test_smallest_of_each_group_then_smallest_of_all(self):
        self.assertEqual(testmode._smallest([[1, 2, 3], [10], []], 4), [1, 2, 3, 10])
        self.assertEqual(testmode._smallest([[1, 2], [10, 20]], 3), [1, 2, 10])
        # a group of big tables adds only its smallest one, not all of them
        self.assertEqual(testmode._smallest([[1, 2, 3, 4, 5], [58, 123, 319158]], 4), [1, 2, 3, 58])
        self.assertEqual(testmode._smallest([], 5), [])

    def test_off_by_default(self):
        with mock.patch.object(settings, "TEST_TABLES", 0):
            self.assertFalse(testmode.on())
            self.assertIsNone(testmode.info(65, 44))
        with mock.patch.object(settings, "TEST_TABLES", 5):
            self.assertEqual(testmode.info(65, 44), {"tables": 5, "mappings": 5, "of_tables": 65, "of_mappings": 44})
            self.assertEqual(testmode.info(3, 2)["tables"], 3)

    def test_pick_smallest_with_records(self):
        from app import main
        from app.ATNM import api as atnm_api
        from app.ATNM import required as atnm_required
        from app.ATNM import settings as atnm_settings

        def tab(rows):
            return {"rows": rows}
        pairs = [mock.Mock(id="1"), mock.Mock(id="2")]
        cats = {("1", "source"): {"tables": {"a.big": tab(900), "a.small": tab(5), "a.empty": tab(0), "a.lost": tab(3)}},
                ("1", "target"): {"tables": {"a.big": tab(900), "a.small": tab(5), "a.empty": tab(0)}},
                ("2", "source"): {"tables": {"b.mid": tab(50), "b.tiny": tab(1)}},
                ("2", "target"): {"tables": {"b.mid": tab(50), "b.tiny": tab(2)}}}
        req = {"1": {k: {} for k in ("a.big", "a.small", "a.empty", "a.lost")}, "2": {"b.mid": {}, "b.tiny": {}}}
        live = {"s1": {"rows": 10}, "t1": {"rows": 12}, "s2": {"rows": 3}, "t2": {"rows": 3},
                "s3": {"rows": 0}, "t3": {"rows": 0}, "s4": {"rows": 1, "locked": True}, "t4": {"rows": 1},
                "s5": {"rows": 7}, "t5": None}
        maps = []
        for i in range(1, 6):
            s, t = mock.Mock(), mock.Mock()
            s.ref, t.ref = f"s{i}", f"t{i}"
            maps.append(mock.Mock(id=f"m{i}", sources=[s], targets=[t]))
        with mock.patch.object(settings, "TEST_TABLES", 3), \
                mock.patch.object(atnm_settings, "PAIRS", pairs), \
                mock.patch.object(atnm_api, "_read_all", return_value=cats), \
                mock.patch.object(atnm_required, "tables", side_effect=lambda p: req[p.id]), \
                mock.patch.object(main, "get_plan", return_value=mock.Mock(mappings=maps)), \
                mock.patch.object(main, "_live_tables", return_value={}), \
                mock.patch.object(main, "_entry", side_effect=lambda lv, r: live[r]):
            sel = testmode.pick()
        # Section 1: the smallest of each pair, then the smallest of all; no empty table, none missing on a side
        self.assertEqual(sel["part1"], {"1": ["a.small"], "2": ["b.tiny", "b.mid"]})
        # Part 2: no empty, locked or missing table; smallest first
        self.assertEqual(sel["part2"], ["m2", "m1"])
        self.assertEqual((sel["of_tables"], sel["of_mappings"]), (6, 5))
        self.assertIn("3 of 6 required tables", sel["text"])
        self.assertIn("2 of 5 mappings", sel["text"])
        self.assertEqual(testmode.keep1(sel, "2"), {"b.tiny", "b.mid"})
        self.assertIsNone(testmode.keep1(None, "2"))
        self.assertEqual(testmode.keep2(sel), {"m2", "m1"})

    def test_test_run_is_marked_and_counts_the_picked_ones(self):
        sel = {"part1": {"1": ["dbo.t0", "dbo.t1"]}, "part2": ["m0", "m1"], "of_tables": 65, "of_mappings": 44}
        sel["text"] = testmode.text(sel)
        with tempfile.TemporaryDirectory() as d:
            ev = evidence(("Correct", "Correct"), ("Correct", "Correct"))
            pf = {"ok": True, "atnm": {"ok": True}, "rds": {"ok": True}, "busy": [],
                  "plan": {"ok": True, "tables": 65, "mappings": 44}, "ai": {"enabled": False}}
            seen = {}

            def part1(r, s):
                seen["part1"] = jobs.job["test"]
                jobs._stage("part1", "done", done=2, total=2)

            with mock.patch.object(settings, "TEST_TABLES", 2), \
                    mock.patch.object(settings, "REPORT_DIR", Path(d)), \
                    mock.patch.object(settings, "RUN_FILE", Path(d) / "run.json"), \
                    mock.patch.object(jobs, "preflight", return_value=pf), \
                    mock.patch.object(testmode, "pick", return_value=sel), \
                    mock.patch.object(jobs, "_part1", side_effect=part1), \
                    mock.patch.object(jobs, "_part2", side_effect=lambda r, s: jobs._stage("part2", "done", done=2, total=2)), \
                    mock.patch("app.report.collect.build", side_effect=lambda run, since: {**ev, "run": {**ev["run"], **run}}), \
                    mock.patch.object(synth, "_write", side_effect=lambda e, n: {"source": "template", **synth.template(e)}):
                ok, msg = jobs.start()
                self.assertTrue(ok, msg)
                for _ in range(100):
                    if not jobs.job["running"]:
                        break
                    time.sleep(0.1)
                st = jobs.status(0)
                self.assertIsNone(st["error"], st["error"])
                self.assertIs(seen["part1"], sel)
                meta = st["result"]
                self.assertEqual(meta["self_check"], [])      # 2 + 2 expected, not 65 + 44
                self.assertEqual(meta["test"], sel["text"])
                self.assertTrue(meta["files"]["docx"].endswith("_TEST.docx"))
                self.assertTrue(meta["files"]["xlsx"].endswith("_TEST.xlsx"))
                self.assertTrue(meta["mode"].startswith(testmode.LABEL))
                from docx import Document
                doc = Document(jobs.file_of(meta["run_id"], "docx"))
                self.assertIn(sel["text"], "\n".join(p.text for p in doc.paragraphs))
                from openpyxl import load_workbook
                self.assertEqual(load_workbook(jobs.file_of(meta["run_id"], "xlsx"))["Summary"]["A2"].value, sel["text"])
                self.assertEqual(jobs._read_run()["test"], sel)   # a resume checks the same ones

    def test_summary_opens_with_the_test_line(self):
        ev = evidence()
        sel = {"part1": {"1": ["dbo.t0"]}, "part2": ["m0"], "of_tables": 65, "of_mappings": 44}
        sel["text"] = testmode.text(sel)
        ev["run"]["test"] = sel
        self.assertIn("run_type: TEST RUN", synth.facts(ev))
        with mock.patch.object(synth, "_write", return_value={"source": "ai", "executive_summary": "All good."}):
            self.assertTrue(synth.write(ev)["executive_summary"].startswith(sel["text"]))
        ev["run"].pop("test")
        self.assertNotIn("run_type", synth.facts(ev))
        with mock.patch.object(synth, "_write", return_value={"source": "ai", "executive_summary": "All good."}):
            self.assertEqual(synth.write(ev)["executive_summary"], "All good.")


if __name__ == "__main__":
    unittest.main()
