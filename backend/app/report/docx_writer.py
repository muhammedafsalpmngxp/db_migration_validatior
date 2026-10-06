"""The Word report of a run, built only from its evidence: short, in plain words, for a
reader who has not worked on the migration. Cover, summary, then the two sections - 1: the
copy ATNM → RDS, 2: the move RDS → new system - each with every table (or mapping), what must
be fixed and what needs a decision; then what was not checked, sign-off and the words used.
Every detail is in the Excel workbook."""
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import plain, settings
from .xlsx_writer import P1_TABLES, P2_MAPPINGS, local

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)
TONE = {plain.CORRECT: ("C6EFCE", "006100"), plain.MUST_FIX: ("FFC7CE", "9C0006"), plain.DECIDE: ("FFEB9C", "9C5700"),
        plain.NOT_CHECKED: ("D9D9D9", "3F3F3F"), plain.BY_DESIGN: ("DDEBF7", "1F4E78"),
        "NOT READY": ("FFC7CE", "9C0006"), "INCOMPLETE": ("FFEB9C", "9C5700"), "READY": ("C6EFCE", "006100")}
WIDTH = 17.0    # cm of text width (A4, 2 cm margins)


def _n(x):
    return f"{x:,}" if isinstance(x, int) and not isinstance(x, bool) else ("—" if x is None else str(x))


def _shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _field(paragraph, code, size):
    run = paragraph.add_run()
    run.font.size = Pt(size)
    for kind, text in (("begin", None), (None, code), ("separate", None), ("text", "1"), ("end", None)):
        if kind in ("begin", "separate", "end"):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = text
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {code} "
        run._r.append(el)


class Doc:
    def __init__(self):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
        st = self.d.styles["Normal"]
        st.font.name, st.font.size = "Arial", Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        for name, size in (("Heading 1", 15), ("Heading 2", 12)):
            h = self.d.styles[name]
            h.font.name, h.font.size, h.font.color.rgb = "Arial", Pt(size), NAVY
            rf = h.element.rPr.rFonts
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rf.attrib.pop(qn(a), None)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), "Arial")

    def p(self, text="", size=10, bold=False, color=None, italic=False, align=None, after=6):
        par = self.d.add_paragraph()
        if text:
            r = par.add_run(text)
            r.font.size, r.bold, r.italic = Pt(size), bold, italic
            if color:
                r.font.color.rgb = color
        par.paragraph_format.space_after = Pt(after)
        if align:
            par.alignment = align
        return par

    def h(self, text, level=1):
        self.d.add_heading(text, level)

    def band(self, text, fill):
        """A section title: a heading (for the navigation pane) in white on a coloured band."""
        h = self.d.add_heading("", 1)
        r = h.add_run(text)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        pr = h._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        pr.append(shd)

    def bullets(self, lines, size=10):
        for line in lines:
            par = self.d.add_paragraph(style="List Bullet")
            par.add_run(line).font.size = Pt(size)
            par.paragraph_format.space_after = Pt(2)

    def table(self, headers, rows, widths, result_col=None, size=8.5, header=True):
        t = self.d.add_table(rows=1 if header else 0, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if header:
            for i, h in enumerate(headers):
                c = t.rows[0].cells[i]
                c.text = ""
                r = c.paragraphs[0].add_run(h)
                r.bold, r.font.size = True, Pt(size)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                _shade(c, "1F3864")
            pr = t.rows[0]._tr.get_or_add_trPr()
            el = OxmlElement("w:tblHeader")
            el.set(qn("w:val"), "true")
            pr.append(el)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(_n(v))
                r.font.size = Pt(size)
                if isinstance(v, int) and not isinstance(v, bool):
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if result_col is not None and i == result_col and str(v) in TONE:
                    fill, color = TONE[str(v)]
                    _shade(cells[i], fill)
                    r.bold = True
                    r.font.color.rgb = RGBColor.from_string(color)
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def box(self, title, lines, note=None):
        t = self.d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        c = t.rows[0].cells[0]
        _shade(c, "EEF3FA")
        c.width = Cm(WIDTH)
        par = c.paragraphs[0]
        r = par.add_run(title)
        r.bold, r.font.size = True, Pt(10.5)
        r.font.color.rgb = NAVY
        if note:
            r = par.add_run(f"  ({note})")
            r.italic, r.font.size = True, Pt(8)
            r.font.color.rgb = GREY
        for line in lines:
            q = c.add_paragraph()
            q.add_run(line).font.size = Pt(10)
            q.paragraph_format.space_after = Pt(3)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _where(i, part):
    if part == 1:
        return f"{i['source_db']} · {i['table']}"
    return f"{i['mapping']}\n{i['table']} → {i['target_table']}"


def _problem_rows(issues, part):
    return [[i["id"], _where(i, part) + (f"\nColumn: {i['column']}" if i["column"] else ""), i["finding"], i["action"]]
            for i in issues]


def _by_result(items):
    return sorted(items, key=lambda x: plain.ORDER.get(x["result"], len(plain.ORDER)))


def build(ev, path):
    doc = Doc()
    d = doc.d
    run, p1, p2, t = ev["run"], ev["part1"], ev["part2"], ev["totals"]
    s = ev.get("summary") or {}
    ai_note = "written by AI from the facts and checked" if s.get("source") == "ai" else "plain sentences from the facts"
    tdb = next((x["name"] for x in p2.get("databases") or [] if x["role"] == "target"), "the new system")
    srcs = ", ".join(x["name"] for x in p2.get("databases") or [] if x["role"] != "target")
    test = (run.get("test") or {}).get("text")

    sec = d.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = f"{settings.TITLE} · {run.get('id')}{' · TEST RUN' if test else ''}"
    hp.runs[0].font.size, hp.runs[0].font.color.rgb = Pt(8), GREY
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Confidential · Page ").font.size = Pt(8)
    _field(fp, "PAGE", 8)
    fp.add_run(" of ").font.size = Pt(8)
    _field(fp, "NUMPAGES", 8)

    # ---- cover ----
    doc.p(after=48)
    doc.p(settings.TITLE, 26, True, NAVY, after=4)
    doc.p(f"ATNM → RDS → {tdb}", 14, color=GREY, after=28 if not test else 10)
    if test:
        doc.p(test, 12, True, RGBColor(0x9C, 0x00, 0x06), after=18)
    box = d.add_table(rows=1, cols=1)
    c = box.rows[0].cells[0]
    fill, color = TONE[ev["overall"]]
    _shade(c, fill)
    par = c.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = par.add_run(f"Overall result: {ev['overall']}")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(18), RGBColor.from_string(color)
    q = c.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.add_run(plain.OVERALL_HELP[ev["overall"]]).font.size = Pt(10)
    q = c.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.add_run(f"{t['issues'][plain.MUST_FIX]} must be fixed · {t['issues'][plain.DECIDE]} need a decision · "
              f"{len(ev['not_checked'])} not checked").font.size = Pt(10)
    doc.p(after=20)
    info = []
    if settings.CLIENT:
        info.append(["Client", settings.CLIENT])
    if settings.PREPARED_BY:
        info.append(["Prepared by", settings.PREPARED_BY])
    info += [["Report date", datetime.now().strftime("%d %B %Y")],
             ["Checked", f"{local(run.get('started_at'))} → {local(run.get('finished_at'))}"],
             ["Section 1 · ATNM → RDS", "\n".join(f"{db['source_db']} (ATNM) → {db['target_db']} (RDS): {db['required']} "
                                                   "tables" for db in p1.get("databases") or []) or "—"],
             [f"Section 2 · RDS → {tdb}", f"{srcs} (RDS) → {tdb}: {len(p2['items'])} mappings"]]
    doc.table(["", ""], info, [5.6, 11.4], size=9.5, header=False)
    doc.p("Read-only check: nothing was changed in any database. No passwords or server addresses are in this report. "
          "Every detail is in the Excel file of the same name.", 8.5, italic=True, color=GREY)
    doc.page_break()

    # ---- summary ----
    doc.h("Summary")
    doc.box("In short", [s.get("executive_summary", "")], ai_note)
    order = (plain.CORRECT, plain.MUST_FIX, plain.DECIDE, plain.NOT_CHECKED, plain.BY_DESIGN)
    doc.table(["", "Correct", "Must fix", "Needs a decision", "Not checked", "By design", "Total"], [
        ["Section 1 · ATNM → RDS (tables)"] + [t["part1"][k] for k in order] + [sum(t["part1"].values())],
        [f"Section 2 · RDS → {tdb} (mappings)"] + [t["part2"][k] for k in order] + [sum(t["part2"].values())],
    ], [5.2, 1.8, 1.8, 2.3, 2.1, 2.0, 1.8], size=9)
    if s.get("top_risks"):
        doc.h("Biggest risks", 2)
        doc.bullets(s["top_risks"])
    if s.get("next_steps"):
        doc.h("Next steps", 2)
        doc.bullets([f"{n}. {x}" for n, x in enumerate(s["next_steps"], 1)])
    doc.h("What the results mean", 2)
    doc.table(["Result", "Meaning"], [[k, v] for k, v in plain.RESULT_HELP.items()], [3.5, 13.5], result_col=0, size=9)
    doc.page_break()

    # ---- the two sections ----
    sections = (
        (1, "part1", "Section 1 · ATNM → RDS: copy of the client's databases", "1F4E79", "part1_summary",
         "The client's databases were copied from the ATNM server to the RDS server; the copy must be exact. For every "
         "table the migration uses, this checks the data: that the table is on both servers, its column names and "
         "data types, the number of records, every value of every record, empty values and renamed columns.",
         ["ATNM database", "RDS copy", "Tables checked"],
         [[db["source_db"], db["target_db"], db["required"]] for db in p1.get("databases") or []],
         "Result of every table", ["Database · table", "Records in ATNM", "Records in RDS", "What was found", "Result"],
         lambda i: [f"{i['source_db']} · {i['table']}" + (f" (RDS: {i['target_table']})" if i.get("target_table")
                                                           and i["target_table"] != i["table"] else ""),
                    i["rows_source"], i["rows_target"], i["reason"], i["result"]],
         P1_TABLES),
        (2, "part2", f"Section 2 · RDS → {tdb}: move into the new system", "385723", "part2_summary",
         f"The copies on RDS were moved into the new system ({tdb}) as the migration plan says: tables renamed, reshaped "
         "and linked. For every mapping, this checks the number of records, how each source column maps to a target "
         "column and its data type, every value of every record, empty values, the links to other tables (do they "
         "point at the right record), and records stored twice or pointing at nothing.",
         ["Source databases (RDS)", "Target database", "Mappings checked"], [[srcs, tdb, len(p2["items"])]],
         "Result of every mapping", ["Mapping: source table → target table", "Records expected", "Records found",
                                     "What was found", "Result"],
         lambda i: [f"{i['mapping']}\n{i['source']} → {i['target']}", i["rows_expected"], i["rows_actual"], i["reason"],
                    i["result"]],
         P2_MAPPINGS),
    )
    for part, key, title, color, summary_key, intro, db_heads, db_rows, list_title, heads, row, prefix in sections:
        items = ev[key]["items"]
        doc.band(title, color)
        doc.p(intro, 9.5, color=GREY)
        doc.table(db_heads, db_rows, [6.0, 6.0, 5.0], size=9)
        doc.box("Summary", [s.get(summary_key, "")], ai_note)
        doc.h(f"{list_title} ({len(items)})", 2)
        doc.table(heads, [row(i) for i in _by_result(items)], [5.0, 2.0, 2.0, 5.6, 2.4], result_col=4, size=8)
        must = [i for i in ev["issues"] if i["part"] == part and i["result"] == plain.MUST_FIX]
        decide = [i for i in ev["issues"] if i["part"] == part and i["result"] == plain.DECIDE]
        for label, rows in ((f"Must fix ({len(must)})", must), (f"Needs a decision ({len(decide)})", decide)):
            doc.h(label, 2)
            if rows:
                doc.table(["Ref", "Where", "What was found", "What to do"], _problem_rows(rows, part),
                          [1.3, 4.6, 6.4, 4.7], size=8)
            else:
                doc.p("Nothing.", 10)
        renames = [x for x in ev[key]["renames"] if x["verdict"] == "verified"]
        if renames:
            doc.h(f"Columns with a new name, proven by their data ({len(renames)})", 2)
            doc.table(["Where", "Source column", "Target column", "Records compared"],
                      [[x.get("mapping") or f"{x['database']} · {x['table']}", x["source"], x["target"],
                        x.get("paired") or x.get("rows_checked")] for x in renames], [5.6, 4.0, 4.6, 2.8], size=8)
        doc.p(f"Every detail of this section (columns, renames, empty values, records) is in the Excel file, on the "
              f"sheet \"{prefix}\".", 8.5, italic=True, color=GREY)
        doc.page_break()

    # ---- not checked, sign-off, words ----
    doc.h("What could not be checked")
    if ev["not_checked"]:
        doc.table(["Section", "Where", "What", "Why", "What to do"],
                  [[n["part"], n["item"], n["what"], n["reason"], n["todo"]] for n in ev["not_checked"]],
                  [1.4, 3.8, 3.0, 5.4, 3.4], size=8)
    else:
        doc.p("Everything was checked.", 10)
    doc.h("Sign-off")
    doc.table(["Role", "Name", "Decision", "Date", "Signature"],
              [["Migration team", "", "", "", ""], ["DB team", "", "", "", ""], ["Client", "", "", "", ""]],
              [3.4, 4.0, 3.4, 2.4, 3.8], size=10)
    doc.h("Words used")
    doc.table(["Word", "Meaning"], [list(g) for g in plain.GLOSSARY], [4.0, 13.0], size=8.5)
    d.save(path)
