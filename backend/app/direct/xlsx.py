"""The direct report's Excel workbook, built only from its evidence (no database is read).

    Summary                 the whole report in one short paragraph
    Client → New System     one row per mapping: client tables (ATNM) and their records, new
                            tables (AlTasnimBI) and their records, the rule, columns, status, issues
    Columns                 every column pair of every mapping
    Tables Not Used         tables of each database that no mapping uses
    Words Used              what the words mean
    Full Details            the whole report in words

Uses the migration report's sheet helpers and words (report.xlsx_writer, report.plain), so
both reports look the same; nothing there is changed.
"""
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.properties import PageSetupProperties

from ..report import plain
from ..report.xlsx_writer import DIFF, FONT, GRID, NAVY, OTHER, TONE, S, _and, _link, _put, _sheet, local
from . import settings

SUMMARY = "Summary"
MAPPINGS = "Client → New System"
COLUMNS = "Columns"
NOT_USED = "Tables Not Used"
WORDS = "Words Used"
FULL = "Full Details"
MAIN = {"tab": "7030A0", "head": "5B2C6F", "name": "Client → New System"}
RESULT_COL = "J"
TYPE_WORDS = {"one_to_one": "one to one", "union": "union (stacked)", "merge": "merge (joined)",
              "transform": "transform", "excluded": "excluded"}
COLUMN_WORDS = {"match": "Same", "renamed": "Renamed", "changed": "Changed", "source_only": "Not migrated",
                "target_only": "Only in target"}


def _count(items, *results):
    return sum(1 for i in items if i["result"] in results)


def _tables(members):
    names = "\n".join(f"{x['database'] or '?'} · {x['table']}" + ("" if x["exists"] is not False else " (does not exist)")
                      + (f" ({x['role']})" if x.get("role") else "") for x in members)
    rows = [x["rows"] for x in members]
    if len(members) == 1:
        return names, rows[0]
    return names, "\n".join("—" if r is None else f"{r:,}" for r in rows)


COLUMN_GROUPS = (("match", "same"), ("renamed", "renamed"), ("changed", "changed"),
                 ("source_only", "not migrated"), ("target_only", "new in target"))


def _column_name(r):
    """One column pair in a few words: Name, Old → new, Name (int → bigint)."""
    s, t = r["source"], r["target"]
    if r["status"] == "target_only":
        return t
    if r["status"] == "source_only":
        return s
    name = s if s == t else f"{s} → {t}"
    if "type" in r["diffs"]:
        name += f" ({r['source_type']} → {r['target_type']})"
    elif r["diffs"]:
        name += f" ({', '.join(r['diffs'])})"
    return name


def _columns_text(i):
    """The columns of a mapping: how many, then each group with its column names, one line each."""
    c, rows = i["columns"], i["column_rows"]
    if not c:
        return "—"
    lines = [f"{c.get('source_columns', 0)} client columns:"]
    for status, word in COLUMN_GROUPS:
        names = list(dict.fromkeys(_column_name(r) for r in rows if r["status"] == status))
        if names:
            lines.append(f"{len(names)} {word}: {', '.join(names)}")
    return "\n".join(lines)


def _lines(i):
    found = [f"{x['id']} {S(x['result'])}: {x['finding']}" for x in i["issues"]]
    found += [f"Not checked: {n['what']} - {n['reason']}" for n in i["not_checked"]]
    found += [f"Note: {n}" for n in i["notes"]]
    if not i["issues"] and not i["not_checked"]:
        found.insert(0, "No issue: the tables exist, the record counts follow the rule and every client column has "
                        "a column in the new table." if i["result"] == plain.CORRECT else "No issue.")
    todo = []
    for x in i["issues"]:
        if x["action"] not in todo:
            todo.append(x["action"])
    for n in i["not_checked"]:
        if n["todo"] not in todo:
            todo.append(n["todo"])
    return "\n".join(found), ("\n".join(f"{k}. {t}" for k, t in enumerate(todo, 1)) if len(todo) > 1
                              else (todo[0] if todo else "—"))


def _main(wb, ev):
    rows = []
    for i in sorted(ev["items"], key=lambda x: plain.ORDER.get(x["result"], 9)):
        src, src_rows = _tables(i["sources"])
        tgt, tgt_rows = _tables(i["targets"])
        found, todo = _lines(i)
        rows.append([i["mapping"], TYPE_WORDS.get(i["type"], i["type"]), src, src_rows, tgt, tgt_rows,
                     i["rule"] or "—", i["difference"], _columns_text(i), S(i["result"]), found, todo])
    _sheet(wb, MAPPINGS, MAIN, f"Every mapping of the migration plan: the client's original tables on the ATNM server "
                              f"→ the new tables in {ev['target_db']}. Problems first; the issue numbers match Full Details.",
           [("Mapping", 24), ("Type", 14), ("Client table(s) (ATNM)", 36), ("Records", 11),
            (f"New table(s) ({ev['target_db']})", 34), ("Records", 11), ("Rule", 30), ("Difference", 11),
            ("Columns", 55), ("Status", 15), ("Issues found", 70), ("What to do", 42)],
           rows, result_cols=(10,), wrap=(3, 4, 5, 6, 7, 9, 11, 12), number_formats={8: DIFF},
           title=f"Client (ATNM) → {ev['target_db']}: direct check of every mapping", row_result=10)


def _others(wb, ev):
    rows = []
    for i in ev["items"]:
        for c in i["column_rows"]:
            rows.append([i["mapping"], c["source_table"], c["source"] or "—", c["source_type"] or "—",
                         c["target"] or "—", c["target_type"] or "—", c["match"] or "—",
                         COLUMN_WORDS.get(c["status"], c["status"])])
    _sheet(wb, COLUMNS, OTHER, "Every column of every mapping: the client column, the column it became in the new "
                               "table, and how they were paired (declared in the plan, by name, or by naming convention).",
           [("Mapping", 24), ("Client table", 30), ("Client column", 26), ("Client type", 16), ("New column", 26),
            ("New type", 16), ("Paired by", 12), ("Column", 14)], rows, title="Columns")
    unused = [[t["server"], t["database"], u["table"], u["rows"]] for t in ev["tables"] for u in t["not_used"]]
    _sheet(wb, NOT_USED, OTHER, "For information: tables that no mapping uses. Confirm none was forgotten.",
           [("Server", 16), ("Database", 24), ("Table", 44), ("Records", 14)], unused, title="Tables not used")
    _sheet(wb, WORDS, OTHER, "What the words in this report mean.", [("Word", 28), ("Meaning", 110)],
           [list(w) for w in plain.GLOSSARY[:7]] + [[S(k), v] for k, v in plain.RESULT_HELP.items()], wrap=(2,),
           title="Words used")


def _title(ws, ev, text):
    run = ev["run"]
    _put(ws, 1, 1, text, 18, True, color=NAVY)
    if run.get("test"):
        _put(ws, 2, 1, run["test"], 12, True, color="9C0006")
    _put(ws, 3, 1, f"Checked {local(run.get('started_at'))} → {local(run.get('finished_at'))}  ·  client server (ATNM) "
                   f"→ new system ({ev['target_db']})", italic=True, color="595959")


def summary_text(ev):
    """The whole direct report in one paragraph, from the evidence."""
    items = ev["items"]
    must, decide, nc = _count(items, plain.MUST_FIX), _count(items, plain.DECIDE), _count(items, plain.NOT_CHECKED)
    clients = [d["database"] for d in ev["databases"] if d["database"] != ev["target_db"]]
    out = [f"This report checks the migration straight from the client's original databases on the ATNM server "
           f"({_and(clients)}) to the new system, {ev['target_db']}, mapping by mapping, without going through the RDS "
           "copies. This version compares the tables, the record counts and the columns; the values themselves are "
           "not compared yet."]

    # What was read
    read = [f"{t['database']} ({t['tables']:,} tables, {t['used']:,} used by the migration)" for t in ev["tables"]
            if t["tables"] is not None]
    unread = [t["database"] for t in ev["tables"] if t["tables"] is None]
    if read:
        out.append(f"We read the table lists of {_and(read)}.")
    if unread:
        out.append(f"{_and(unread)} could not be read in this run, so the mappings that use them are not checked.")

    # The result
    if ev["overall"] == "READY" and decide:
        out.append(f"Nothing that was checked needs to be fixed, but {decide} mapping{'s' if decide != 1 else ''} "
                   f"need{'s' if decide == 1 else ''} a decision from the client before go-live.")
    else:
        out.append(plain.STATUS_SENTENCE.get(ev["overall"], ""))
    rest = [f"{must} need to be fixed"] if must else []
    if decide:
        rest.append(f"{decide} need a decision")
    if nc:
        rest.append(f"{nc} could not be checked")
    out.append(f"We checked {len(items)} mapping{'s' if len(items) != 1 else ''}: "
               f"{_count(items, plain.CORRECT, plain.BY_DESIGN)} are correct" + (", " + _and(rest) if rest else "") + ".")

    # Tables
    missing = [f"{x['table']} ({i['mapping']})" for i in items for x in i["sources"] + i["targets"] if x["exists"] is False]
    if missing:
        out.append(f"{len(missing)} table{'s' if len(missing) != 1 else ''} of the plan "
                   f"{'are' if len(missing) != 1 else 'is'} missing: {_and(missing[:5])}.")
    else:
        out.append("Every table the plan names exists on both servers.")

    # Record counts
    compared = [i for i in items if i["difference"] is not None]
    differ = sorted((i for i in compared if i["difference"]), key=lambda i: -abs(i["difference"]))
    if compared:
        line = (f"The record counts follow the plan's rule in {len(compared) - len(differ)} of the {len(compared)} "
                "mappings that could be counted")
        if differ:
            line += ("; they differ in " + str(len(differ)) + ", the most in "
                     + _and([f"{i['mapping']} ({i['difference']:+,})" for i in differ[:3]]))
        out.append(line + ".")

    # Columns
    groups = {s: 0 for s, _ in COLUMN_GROUPS}
    for i in items:
        for r in i["column_rows"]:
            groups[r["status"]] = groups.get(r["status"], 0) + 1
    total = groups["match"] + groups["renamed"] + groups["changed"] + groups["source_only"]
    if total:
        out.append(f"Across these mappings, {total:,} client columns were compared: {groups['match']:,} kept the same "
                   f"name and type, {groups['renamed']:,} were renamed, {groups['changed']:,} changed type or settings, "
                   f"and {groups['source_only']:,} have no column in the new table; the new tables also have "
                   f"{groups['target_only']:,} columns of their own.")
    lost = [f"{i['mapping']} ({', '.join(r['source'] for r in i['column_rows'] if r['status'] == 'source_only')})"
            for i in items if any(r["status"] == "source_only" for r in i["column_rows"])]
    if lost:
        out.append(f"The client must confirm the columns that were not carried over: {_and(lost[:5])}"
                   + (f" and {len(lost) - 5} more mappings" if len(lost) > 5 else "") + ".")
    out.append("Every mapping, with its column names, what is wrong and what to do, is listed in the other sheets of "
               "this report.")
    return " ".join(x for x in out if x)


def _summary(wb, ev):
    ws = wb[SUMMARY]
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 125
    _title(ws, ev, settings.TITLE)
    text = summary_text(ev)
    _put(ws, 5, 1, text, 11, wrap=True)
    ws.row_dimensions[5].height = 16 * max(1, -(-len(text) // 115)) + 8
    _put(ws, 7, 1, f"Run {ev['run']['id']}. Read-only check of the catalogs: nothing was changed in any database. No "
                   "passwords or server addresses are in this file.", 8, italic=True, color="7F7F7F")


def _full(wb, ev):
    ws = wb.create_sheet(FULL)
    ws.sheet_properties.tabColor = "7030A0"
    ws.sheet_view.showGridLines = False
    for c, w in zip("ABCDEF", (28, 26, 16, 18, 16, 34)):
        ws.column_dimensions[c].width = w
    _title(ws, ev, "Full details of the direct report")
    items = ev["items"]
    r = 5

    def title(t):
        nonlocal r
        _put(ws, r, 1, t, 13, True, color=NAVY)
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = Border(bottom=Side(style="medium", color=NAVY))
        r += 1

    def text(t):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        _put(ws, r, 1, t, wrap=True)
        ws.row_dimensions[r].height = 14 * max(1, -(-len(t) // 125)) + 6
        r += 1

    def header(names):
        nonlocal r
        for c, name in enumerate(names, 1):
            cell = _put(ws, r, c, name, bold=True, color="FFFFFF", fill=NAVY, wrap=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = GRID
        r += 1

    def cells(values, centered=(3, 4, 5)):
        nonlocal r
        for c, v in enumerate(values, 1):
            cell = _put(ws, r, c, "—" if v is None else v, wrap=True)
            cell.border = GRID
            if c in centered:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            if isinstance(v, int) and not isinstance(v, bool):
                cell.number_format = "#,##0"
        r += 1

    title("1. About this report")
    text(f"This report checks the migration directly from the client's original databases on the ATNM server to the new "
         f"system ({ev['target_db']}), using the mappings of the migration plan: one client table to one new table, "
         "several client tables stacked into one (union), or joined into one (merge). It does not use the RDS copies. "
         "This version reads the table lists only (names, columns, types and record counts); the values are not "
         "compared yet. Nothing was changed in any database.")
    r += 1

    title("2. Overall result")
    col = f"'{MAPPINGS}'!{RESULT_COL}:{RESULT_COL}"
    must, nc = S(plain.MUST_FIX), S(plain.NOT_CHECKED)
    _put(ws, r, 1, "Status", 11, True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    _put(ws, r, 2, f"=IF(COUNTIF({col},\"{must}\")>0,\"NOT READY\",IF(COUNTIF({col},\"{nc}\")>0,\"INCOMPLETE\",\"READY\"))",
         12, True)
    for word in ("NOT READY", "INCOMPLETE", "READY"):
        fill, color = TONE[word]
        ws.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f'B{r}="{word}"'], fill=PatternFill("solid", fgColor=fill),
                                                           font=Font(name=FONT, size=12, bold=True, color=color)))
    r += 1
    text(f"{plain.STATUS_SENTENCE.get(ev['overall'], '')} {_count(items, plain.MUST_FIX)} mappings need to be fixed, "
         f"{_count(items, plain.DECIDE)} need a decision and {_count(items, plain.NOT_CHECKED)} could not be checked; "
         f"{_count(items, plain.CORRECT, plain.BY_DESIGN)} of {len(items)} are correct.")
    r += 1

    title("3. Table count")
    header(["Server", "Database", "Tables", "Used by the migration", "Not used", "Note"])
    for t in ev["tables"]:
        note = ("could not be read" if t["error"] else
                "the new system" if t["database"] == ev["target_db"] else "client database (original)")
        cells([t["server"], t["database"], t["tables"], t["used"],
               None if t["tables"] is None else t["tables"] - t["used"], note])
    text(f"The tables no mapping uses are listed in the sheet '{NOT_USED}'.")
    r += 1

    title("4. Results")
    header(["", "Mappings", S(plain.CORRECT), S(plain.MUST_FIX), S(plain.DECIDE), S(plain.NOT_CHECKED)])
    cells(["Client → New System", len(items)] + [f"=COUNTIF({col},\"{S(w)}\")" for w in
                                                 (plain.CORRECT, plain.MUST_FIX, plain.DECIDE, plain.NOT_CHECKED)],
          centered=(2, 3, 4, 5, 6))
    _link(ws, r - 1, 1, MAPPINGS, "Client → New System")
    ws.cell(row=r - 1, column=1).border = GRID
    if _count(items, plain.BY_DESIGN):
        text(f"{_count(items, plain.BY_DESIGN)} mapping(s) are reshaped on purpose (By design: transform or excluded).")
    r += 1

    title("5. Findings")
    kinds = {}
    for x in ev["issues"]:
        k = kinds.setdefault((x["result"], x["check"]), [])
        if x["mapping"] not in k:
            k.append(x["mapping"])
    if kinds:
        for (result, kind), names in sorted(kinds.items(), key=lambda kv: (plain.ORDER[kv[0][0]], -len(kv[1]))):
            short = plain.short(kind)
            text(f"•  {short[:1].upper() + short[1:]} ({S(result)}): {len(names)} mapping"
                 f"{'s' if len(names) != 1 else ''}, e.g. {_and(names[:5])}.")
    else:
        text("Nothing to fix or decide in what this version checks.")
    r += 1

    title("6. Decisions needed from the client")
    asks = [x for x in ev["issues"] if x["result"] == plain.DECIDE]
    for x in asks:
        text(f"•  {x['mapping']}: {x['finding']}")
    if not asks:
        text("No decision is needed.")
    r += 1

    title("7. How the check was done, and its limits")
    text("Every plan table on side A or B was looked up as the client's original table on the ATNM server (the database "
         "whose copy RDS holds, with the plan's atnm_names for renamed copies); side T is the new system. Record counts "
         "come from the catalog (sys.partitions), without counting row by row. Columns were paired with the same rules "
         "as the RDS check: declared in the plan, then by name, then by naming convention (marked as a guess).")
    text("Limits of this version: the values are not compared, so a correct record count does not yet prove the data "
         "is the same. The client server is live, while the new system holds the data of its last load, so record counts "
         "can differ until an agreed cut-off.")
    r += 1

    title("8. What the status means")
    for word in plain.RESULTS:
        fill, color = TONE[S(word)]
        _put(ws, r, 1, S(word), bold=True, color=color, fill=fill).border = GRID
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        _put(ws, r, 2, plain.RESULT_HELP[word], wrap=True)
        r += 1
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def build(ev, path):
    wb = Workbook()
    wb.active.title = SUMMARY
    _main(wb, ev)
    _others(wb, ev)
    _full(wb, ev)
    _summary(wb, ev)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(path)
